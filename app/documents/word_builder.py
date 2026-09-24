"""Word-Anlage zur VRA — die eigentliche behördenfertige Datei.

Wichtig — was NIEMALS ins Word darf (Build-Prompt §54):

- LLM-Rohantworten oder Prompt-Ausschnitte
- Confidence-Werte
- Modell-Metadaten (Provider, Version, Token-Kosten)
- Debug- oder Trace-Informationen
- API-Daten
- Das behördliche Anschreiben (macht die Behörde selber)

Was das Word enthält:

- Deckblatt (Projekt, Ort, Zeitraum, Auftragnehmer)
- Übersichtstabelle NVT/Adresse/Regelplan/Status
- Pro NVT ein Abschnitt mit Originalfoto, Absicherungs-Darstellung,
  Regelplan-Preview, Verkehrssituation, Prüfstatus, Begründung
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor

from app.core.enums import DecisionCode, NvtStatus, PhotoKind
from app.core.models import Nvt, Project
from app.ruleplans import RulePlanLibrary


@dataclass
class WordBuildResult:
    docx_path: Path
    included_nvt_ids: list[str] = field(default_factory=list)
    skipped_nvt_ids: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _fmt_date(d: date | None) -> str:
    return d.strftime("%d.%m.%Y") if d else "—"


def _address_line(nvt: Nvt) -> str:
    addr = nvt.address_json or {}
    parts: list[str] = []
    street = addr.get("street") or ""
    hnr = addr.get("house_number") or ""
    line1 = f"{street} {hnr}".strip()
    if line1:
        parts.append(line1)
    plz = addr.get("postal_code") or ""
    city = addr.get("city") or ""
    line2 = f"{plz} {city}".strip()
    if line2:
        parts.append(line2)
    return ", ".join(parts) if parts else "Adresse unbekannt"


def _status_label(nvt: Nvt) -> str:
    s = NvtStatus(nvt.status)
    return {
        NvtStatus.NEW: "neu",
        NvtStatus.ANALYZING: "in Analyse",
        NvtStatus.ANALYZED: "analysiert",
        NvtStatus.NEEDS_REVIEW: "manuelle Prüfung",
        NvtStatus.REVIEWED: "geprüft",
        NvtStatus.APPROVED: "freigegeben",
        NvtStatus.REJECTED: "abgelehnt",
        NvtStatus.EXPORTED: "exportiert",
    }.get(s, s.value)


def _ruleplan_label(nvt: Nvt) -> str:
    if nvt.decision is None:
        return "—"
    if nvt.decision.selected_ruleplan_id:
        return nvt.decision.selected_ruleplan_id
    code = DecisionCode(nvt.decision.code)
    if code == DecisionCode.PRIVATFLAECHE:
        return "Privatfläche"
    return "—"


def _first_photo(nvt: Nvt, kind: PhotoKind):  # type: ignore[no-untyped-def]
    for p in nvt.photos:
        p_kind = p.kind.value if hasattr(p.kind, "value") else str(p.kind)
        if p_kind == kind.value:
            return p
    return None


def _verkehrssituation_text(nvt: Nvt) -> str:
    """Fachliche Kurzbeschreibung, aus Environment abgeleitet."""
    from app.core.enums import Ternary

    if nvt.environment is None:
        return "Keine Umgebungsanalyse vorhanden."

    parts: list[str] = []
    e = nvt.environment
    if Ternary(e.road_present) == Ternary.YES:
        parts.append("Fahrbahn vorhanden")
    if Ternary(e.sidewalk_present) == Ternary.YES:
        parts.append("Gehweg vorhanden")
    if Ternary(e.cycleway_present) == Ternary.YES:
        parts.append("Radweg vorhanden")
    if Ternary(e.intersection_present) == Ternary.YES:
        parts.append("Kreuzung im Nahbereich")
    if Ternary(e.junction_present) == Ternary.YES:
        parts.append("Einmündung im Nahbereich")
    if Ternary(e.cul_de_sac) == Ternary.YES:
        parts.append("Sackgasse")
    if Ternary(e.bus_stop_nearby) == Ternary.YES:
        parts.append("Bushaltestelle in der Nähe")
    if Ternary(e.private_property) == Ternary.YES:
        parts.append("Privatfläche")
    if Ternary(e.business_property) == Ternary.YES:
        parts.append("Betriebsgelände")
    return "; ".join(parts) or "Situation aus vorliegender Analyse nicht spezifisch."


def _reasons_text(nvt: Nvt) -> str:
    """Fachlich formulierte Begründung — OHNE Confidence/Modell/Trace."""
    if nvt.decision is None:
        return "—"
    reasons = list(nvt.decision.reasons or [])
    # Confidence-Werte oder Score-Prosa herausfiltern (Halluzinations-Schutz)
    clean = [r for r in reasons if "Score" not in r and "confidence" not in r.lower()]
    return "  ".join(clean) if clean else "Begründung siehe Prüfhinweise."


def _add_deckblatt(doc: Document, project: Project) -> None:
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("TECHNISCHE UNTERLAGEN")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run("Verkehrssicherung / NVT-Einblasarbeiten")
    sub_run.font.size = Pt(14)

    doc.add_paragraph()  # Abstand

    # Meta-Tabelle
    table = doc.add_table(rows=6, cols=2)
    table.style = "Light Grid Accent 1"
    meta_rows = [
        ("Projekt", project.name),
        ("Ort", project.location_label or "—"),
        ("Zeitraum", f"{_fmt_date(project.period_start)} – {_fmt_date(project.period_end)}"),
        ("Auftraggeber", project.client or "—"),
        ("Auftragnehmer", project.contractor),
        ("Verantwortlich vor Ort",
         f"{project.on_site_responsible_name or '—'}   Tel: {project.on_site_responsible_phone or '—'}"),
    ]
    for i, (label, value) in enumerate(meta_rows):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = str(value)

    doc.add_paragraph()
    hint = doc.add_paragraph(
        "Die nachfolgenden Angaben sind die vom Auftragnehmer erstellten "
        "technischen Anlagen zur verkehrsrechtlichen Anordnung. Das behördliche "
        "Anschreiben wird von der zuständigen Straßenverkehrsbehörde separat "
        "erstellt und ist nicht Bestandteil dieses Dokuments."
    )
    for r in hint.runs:
        r.font.size = Pt(9)
        r.italic = True
    doc.add_page_break()


def _add_overview_table(doc: Document, nvts: list[Nvt]) -> None:
    doc.add_heading("Übersicht", level=1)
    if not nvts:
        doc.add_paragraph("Keine NVT im Export enthalten.")
        return

    table = doc.add_table(rows=1 + len(nvts), cols=4)
    table.style = "Light Grid Accent 1"
    headers = ["NVT", "Adresse", "Regelplan", "Status"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        for r in cell.paragraphs[0].runs:
            r.bold = True

    for row_idx, nvt in enumerate(nvts, start=1):
        table.cell(row_idx, 0).text = nvt.nvt_number
        table.cell(row_idx, 1).text = _address_line(nvt)
        table.cell(row_idx, 2).text = _ruleplan_label(nvt)
        table.cell(row_idx, 3).text = _status_label(nvt)

    doc.add_page_break()


_OVERLAY_WIDTH = Inches(5.5)


def _add_editable_overlay(doc: Document, nvt: Nvt, base_photo: Path) -> list[str]:
    """Originalfoto + Absperrung als bewegliche Word-Zeichenobjekte.

    Statt des fertig gerenderten ``*_proposal.jpg`` wird hier das unveränderte
    Foto eingebettet und die Absperrung darüber als eigenständige Formen
    gelegt, damit der Fachanwender sie in Word verschieben kann.
    """
    from app.core.schemas import VisualizationSchema
    from app.documents.word_overlay import (
        add_movable_overlay,
        photo_placement,
    )

    viz = nvt.visualization
    schema = VisualizationSchema(
        nvt_id=nvt.id,
        base_photo_id=viz.base_photo_id,
        symbols=list(viz.symbols or []),
        shapes=list(viz.shapes or []),
        rendered_photo_id=viz.rendered_photo_id,
        final_photo_id=viz.final_photo_id,
        edited_by_user=viz.edited_by_user,
        edited_at=viz.edited_at,
    )

    para = doc.add_paragraph()
    para.add_run().add_picture(str(base_photo), width=_OVERLAY_WIDTH)
    placement = photo_placement(base_photo, int(_OVERLAY_WIDTH))
    warnings = add_movable_overlay(
        doc, para, visualization=schema, placement=placement,
    )

    hint = doc.add_paragraph(
        "Hinweis: Absperrkante, Baken und Beschriftungen sind bewegliche "
        "Zeichenobjekte — in Word anklicken und verschieben. Das Foto darunter "
        "bleibt unverändert."
    )
    for r in hint.runs:
        r.font.size = Pt(8)
        r.italic = True
    return warnings


def _add_nvt_section(
    doc: Document, nvt: Nvt, library: RulePlanLibrary, *, editable_overlay: bool = False
) -> list[str]:
    warnings: list[str] = []
    doc.add_heading(f"NVT {nvt.nvt_number}", level=1)

    # Kopf-Metadaten
    meta = doc.add_paragraph()
    meta.add_run("Adresse: ").bold = True
    meta.add_run(_address_line(nvt))

    if nvt.location_json and nvt.location_json.get("latitude"):
        gps = doc.add_paragraph()
        gps.add_run("GPS: ").bold = True
        gps.add_run(
            f"{nvt.location_json['latitude']}, {nvt.location_json['longitude']}"
        )

    rp = doc.add_paragraph()
    rp.add_run("Regelplan: ").bold = True
    rp.add_run(_ruleplan_label(nvt))

    st = doc.add_paragraph()
    st.add_run("Prüfstatus: ").bold = True
    st.add_run(_status_label(nvt))

    # Verkehrssituation
    doc.add_heading("Verkehrssituation", level=2)
    doc.add_paragraph(_verkehrssituation_text(nvt))

    # Begründung
    if nvt.decision is not None:
        doc.add_heading("Begründung der Regelplan-Auswahl", level=2)
        doc.add_paragraph(_reasons_text(nvt))

        # Warnungen aus Rule Engine (falls vorhanden)
        if nvt.decision.warnings:
            doc.add_heading("Prüfhinweise", level=3)
            for w in nvt.decision.warnings:
                doc.add_paragraph(w, style="List Bullet")

    # Originalfoto
    original = _first_photo(nvt, PhotoKind.ORIGINAL)
    if original and Path(original.stored_path).is_file():
        doc.add_heading("Originalfoto", level=2)
        doc.add_picture(str(original.stored_path), width=Inches(5.5))
    else:
        warnings.append(f"NVT {nvt.nvt_number}: Originalfoto nicht gefunden")

    # Absicherungs-Darstellung
    proposal = _first_photo(nvt, PhotoKind.FINAL) or _first_photo(nvt, PhotoKind.PROPOSAL)
    if proposal and Path(proposal.stored_path).is_file():
        doc.add_heading("Absicherungs-Darstellung", level=2)
        if editable_overlay and nvt.visualization is not None and original:
            warnings.extend(
                _add_editable_overlay(doc, nvt, Path(original.stored_path))
            )
        else:
            doc.add_picture(str(proposal.stored_path), width=Inches(5.5))
    else:
        code = DecisionCode(nvt.decision.code) if nvt.decision else None
        if code != DecisionCode.PRIVATFLAECHE:
            warnings.append(f"NVT {nvt.nvt_number}: Absicherungs-Darstellung fehlt")

    # Regelplan-Preview
    if nvt.decision and nvt.decision.selected_ruleplan_id:
        entry = library.get(nvt.decision.selected_ruleplan_id)
        if entry:
            preview = Path(entry.schema.preview_png_path)
            if preview.is_file():
                doc.add_heading("Regelplan-Abbildung", level=2)
                doc.add_picture(str(preview), width=Inches(5.5))
                credit = doc.add_paragraph(f"Quelle: {entry.schema.quelle} — Rev. {entry.schema.revision_hash[:8]}")
                for r in credit.runs:
                    r.font.size = Pt(8)
                    r.italic = True

    doc.add_page_break()
    return warnings


def build_word_anlage(
    project: Project,
    nvts: list[Nvt],
    ruleplan_library: RulePlanLibrary,
    output_path: Path,
    *,
    editable_overlay: bool = False,
) -> WordBuildResult:
    """Erzeugt die technische Anlage zur VRA.

    ``editable_overlay=True`` bettet statt des fertig gerenderten
    Absicherungs-Fotos das Originalfoto plus bewegliche Word-Zeichenobjekte
    ein, damit der Fachanwender die Absperrung direkt in Word verschieben
    kann (siehe ``app/documents/word_overlay.py``).
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # Seitenränder etwas enger
    for section in doc.sections:
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    # Neutrale Kern-Properties (keine Modell/Session-Info)
    props = doc.core_properties
    props.title = f"VRA-Anlagen · {project.name}"
    props.author = project.contractor
    props.comments = ""  # bewusst leer

    _add_deckblatt(doc, project)
    _add_overview_table(doc, nvts)

    warnings: list[str] = []
    included: list[str] = []
    skipped: list[str] = []
    for nvt in nvts:
        try:
            warns = _add_nvt_section(
                doc, nvt, ruleplan_library, editable_overlay=editable_overlay
            )
            warnings.extend(warns)
            included.append(str(nvt.id))
        except Exception as exc:  # pragma: no cover — defensiv
            warnings.append(f"NVT {nvt.nvt_number}: Fehler beim Rendern — {exc}")
            skipped.append(str(nvt.id))

    doc.save(str(output_path))
    return WordBuildResult(
        docx_path=output_path,
        included_nvt_ids=included,
        skipped_nvt_ids=skipped,
        warnings=warnings,
    )
