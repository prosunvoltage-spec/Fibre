"""Liest die Excel-Eingabemaske und erzeugt das ausgefuellte PDF-Aufmassblatt.

Aufruf:
    python3 aufmass_pdf.py [Aufmassblatt.xlsx] [Ziel.pdf]

Ohne Argumente wird "Aufmassblatt.xlsx" neben diesem Skript gelesen und das
Ergebnis als "Aufmass_<NVT-Gebiet>_<Datum>.pdf" daneben abgelegt.
"""
import datetime
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

from make_template import KOPF_ZEILEN, KOPF_START, TAB_START
from layout import MAX_ROWS
from pdf_fill import fuellen

HIER = Path(__file__).resolve().parent
VORLAGE = HIER / "vorlage" / "Blanko_Aufmassblatt.pdf"


def _text(wert):
    if wert is None:
        return ""
    if isinstance(wert, (datetime.datetime, datetime.date)):
        return wert.strftime("%d.%m.%Y")
    if isinstance(wert, float) and wert.is_integer():
        return str(int(wert))
    return str(wert).strip()


def _dateiname_teil(text, ersatz):
    sauber = re.sub(r"[^A-Za-z0-9_-]+", "_", text).strip("_")
    return sauber or ersatz


def lesen(xlsx):
    wb = load_workbook(xlsx, data_only=True)
    if "Aufmaß" not in wb.sheetnames:
        raise ValueError(f"{xlsx}: Das Blatt „Aufmaß“ fehlt.")
    ws = wb["Aufmaß"]

    roh = {}
    for i, (_, _, schluessel) in enumerate(KOPF_ZEILEN):
        roh[schluessel] = _text(ws.cell(KOPF_START + i, 2).value)

    ort, datum = roh.pop("ort"), roh.pop("datum")
    roh["ort_datum"] = ", ".join(t for t in (ort, datum) if t)

    zeilen = []
    for n in range(MAX_ROWS):
        z = TAB_START + n
        taetigkeit = _text(ws.cell(z, 3).value)
        if not taetigkeit:
            continue
        zeilen.append({
            "nvt": _text(ws.cell(z, 2).value),
            "taetigkeit": taetigkeit,
            "menge": ws.cell(z, 6).value,
            "bemerkung": _text(ws.cell(z, 7).value),
        })
    if not zeilen:
        raise ValueError(
            f"{xlsx}: Keine Position eingetragen – bitte mindestens eine "
            f"Tätigkeit auswählen.")
    return roh, zeilen, datum


def main(argv):
    xlsx = Path(argv[1]) if len(argv) > 1 else HIER / "Aufmassblatt.xlsx"
    if not xlsx.exists():
        raise SystemExit(f"Eingabedatei nicht gefunden: {xlsx}")
    if not VORLAGE.exists():
        raise SystemExit(f"Blanko-Vorlage nicht gefunden: {VORLAGE}")

    kopf, zeilen, datum = lesen(xlsx)

    if len(argv) > 2:
        ziel = Path(argv[2])
    else:
        gebiet = _dateiname_teil(kopf.get("nvt_gebiet", ""), "NVT")
        tag = _dateiname_teil(datum, datetime.date.today().strftime("%d.%m.%Y"))
        ziel = xlsx.with_name(f"Aufmass_{gebiet}_{tag}.pdf")

    fuellen(str(VORLAGE), str(ziel), kopf, zeilen)
    print(f"Fertig: {ziel}  ({len(zeilen)} Position(en))")
    return ziel


if __name__ == "__main__":
    try:
        main(sys.argv)
    except Exception as fehler:          # Fehler gut lesbar fuer Endnutzer
        print(f"\nFEHLER: {fehler}\n", file=sys.stderr)
        raise SystemExit(1)
