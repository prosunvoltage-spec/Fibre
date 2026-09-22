"""Erzeugt die Excel-Eingabemaske Aufmassblatt.xlsx aus layout.py.

Aufruf:  python3 make_template.py [zieldatei.xlsx]
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from layout import ANZEIGE, MAX_ROWS, TAETIGKEITEN

KOPF_ZEILEN = [
    ("Ort",                  "Münster",              "ort"),
    ("Datum",                "",                     "datum"),
    ("NVT Gebiet",           "",                     "nvt_gebiet"),
    ("Unternehmen",          "Helder Santos GmbH",   "unternehmen"),
    ("BV",                   "Norbert Diez",         "bv"),
    ("Projekt",              "",                     "projekt"),
    ("Projekt (Zusatz)",     "",                     "projekt_zusatz"),
    ("Betriebsnetz",         "",                     "betriebsnetz"),
    ("Betriebsnetz (Zusatz)", "",                    "betriebsnetz_zusatz"),
]

TAB_KOPF = ["Nr.", "NVT", "Tätigkeit", "Position", "Einheit", "Menge",
            "Bemerkung"]
BREITEN = [5, 12, 42, 14, 9, 12, 34]

BLAU = "1F4E79"
HELLBLAU = "DCE6F1"
GRAU = "F2F2F2"
duenn = Side(style="thin", color="B0B0B0")
RAHMEN = Border(left=duenn, right=duenn, top=duenn, bottom=duenn)

KOPF_START = 3
TAB_KOPF_ZEILE = KOPF_START + len(KOPF_ZEILEN) + 2      # 14
TAB_START = TAB_KOPF_ZEILE + 1                          # 15
TAB_ENDE = TAB_START + MAX_ROWS - 1                     # 44


def bauen(ziel):
    wb = Workbook()

    # --- Katalogblatt (Nachschlagetabelle fuer die Dropdowns) ---
    kat = wb.create_sheet("Katalog")
    kat["A1"], kat["B1"], kat["C1"] = "Tätigkeit", "Position", "Einheit"
    for zelle in ("A1", "B1", "C1"):
        kat[zelle].font = Font(bold=True, color="FFFFFF")
        kat[zelle].fill = PatternFill("solid", fgColor=BLAU)
    for i, t in enumerate(TAETIGKEITEN, start=2):
        kat.cell(i, 1, ANZEIGE.get(t["name"], t["name"]))
        kat.cell(i, 2, t["pos"])
        kat.cell(i, 3, t["unit"])
    kat.column_dimensions["A"].width = 42
    kat.column_dimensions["B"].width = 14
    kat.column_dimensions["C"].width = 9
    letzte = len(TAETIGKEITEN) + 1

    # --- Eingabeblatt ---
    ws = wb.active
    ws.title = "Aufmaß"
    ws.sheet_view.showGridLines = False

    ws["A1"] = "Aufmaßblatt HK Hiltrup West / HK Albachten"
    ws["A1"].font = Font(bold=True, size=14, color=BLAU)

    for i, (label, vorgabe, _) in enumerate(KOPF_ZEILEN):
        z = KOPF_START + i
        ws.cell(z, 1, label).font = Font(bold=True)
        ws.cell(z, 1).alignment = Alignment(horizontal="right")
        c = ws.cell(z, 2, vorgabe)
        c.fill = PatternFill("solid", fgColor=HELLBLAU)
        c.border = RAHMEN
    ws.cell(KOPF_START + 1, 2).number_format = "TT.MM.JJJJ"
    ws.cell(KOPF_START + 1, 3, "← Datum eintragen").font = Font(
        italic=True, color="808080")

    # Tabellenkopf
    for sp, titel in enumerate(TAB_KOPF, start=1):
        c = ws.cell(TAB_KOPF_ZEILE, sp, titel)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLAU)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = RAHMEN
        ws.column_dimensions[get_column_letter(sp)].width = BREITEN[sp - 1]

    dv = DataValidation(
        type="list",
        formula1=f"=Katalog!$A$2:$A${letzte}",
        allow_blank=True,
        showDropDown=False,
        errorTitle="Unbekannte Tätigkeit",
        error="Bitte eine Tätigkeit aus der Liste auswählen.",
    )
    ws.add_data_validation(dv)

    for n in range(1, MAX_ROWS + 1):
        z = TAB_START + n - 1
        ws.cell(z, 1, n).alignment = Alignment(horizontal="center")
        ws.cell(z, 1).font = Font(color="808080")

        dv.add(ws.cell(z, 3))
        # Position und Einheit ergeben sich automatisch aus der Tätigkeit
        ws.cell(z, 4, f'=IFERROR(VLOOKUP($C{z},Katalog!$A:$C,2,FALSE),"")')
        ws.cell(z, 5, f'=IFERROR(VLOOKUP($C{z},Katalog!$A:$C,3,FALSE),"")')
        for sp in (4, 5):
            c = ws.cell(z, sp)
            c.fill = PatternFill("solid", fgColor=GRAU)
            c.font = Font(color="606060")
            c.alignment = Alignment(horizontal="center")
        for sp in (2, 6):
            ws.cell(z, sp).fill = PatternFill("solid", fgColor=HELLBLAU)
        ws.cell(z, 6).number_format = "0.##"
        for sp in range(1, len(TAB_KOPF) + 1):
            ws.cell(z, sp).border = RAHMEN

    ws.freeze_panes = ws.cell(TAB_START, 1)

    hinweis = TAB_ENDE + 2
    ws.cell(hinweis, 1,
            "Ausfüllen: NVT, Tätigkeit (Dropdown), Menge, ggf. Bemerkung. "
            "Position und Einheit werden automatisch ergänzt. "
            "Danach „PDF erstellen“ ausführen.")
    ws.cell(hinweis, 1).font = Font(italic=True, color="808080")

    wb.save(ziel)
    return ziel


if __name__ == "__main__":
    ziel = sys.argv[1] if len(sys.argv) > 1 else "Aufmassblatt.xlsx"
    print("erstellt:", bauen(ziel))
